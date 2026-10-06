"""The shop as a spreadsheet (6 Oct 2026, owner's request).

Download writes every product, live or taken down, in exactly the columns
the upload reads — so the file is the audit copy, the template, and the way
to edit forty products at once. Upload reads the same file back.

THE SKU IS THE KEY. A row whose SKU exists updates that product; a new SKU
adds one. A row without a SKU is refused rather than added, because the
same sheet uploaded twice must change nothing the second time — a blank key
would add every such row again.

ALL OR NOTHING. Every row is checked before anything is written; one bad
row and the upload writes nothing, and every problem is listed with its row
number. A half-applied price list is worse than none: nobody can tell which
half went in.

Money is in rupees in the sheet and paise in the database, as everywhere
in the console. GST is a percentage in the sheet and basis points in the
database.
"""

import io
import uuid
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone
from django.utils.text import slugify

from .models import Product, ShopCategory, ShopSubcategory

# (header, help) in sheet order. The header is what the upload looks for;
# the order and any extra columns do not matter.
COLUMNS = (
    ("sku", "Product ID. Required. Unique: an existing SKU updates that product, a new one adds a product."),
    ("name", "Required."),
    ("brand", ""),
    ("subtitle", "One line under the name."),
    ("category", "Required. Must be a category that exists in the console."),
    ("subcategory", "Optional. Must belong to the category."),
    ("price_rupees", "Required. What is charged, in rupees: 1850 not 185000."),
    ("mrp_rupees", "Optional. The struck-through price; must be above the price."),
    ("stock", "Required. Whole number, 0 or more."),
    ("weight_grams", "Required. Shipping is priced on it."),
    ("gst_percent", "0, 3, 5, 12, 18 or 28."),
    ("featured", "yes or no."),
    ("active", "yes puts it in the shop, no takes it down."),
    ("cover_url", "The first photo."),
    ("gallery", "More photos and videos, one URL per line, in order."),
    ("description", "The product page text. New lines are kept."),
    ("faq", "One question per line, written: Question | Answer"),
    ("slug", "The page address. Blank makes one from the name."),
    ("seo_title", "Search and share title. Blank uses the name."),
    ("seo_description", "Search and share line. Blank uses the description's first line."),
    ("id", "Filled by the download. Leave it as it is; never type one."),
)
HEADERS = [c for c, _ in COLUMNS]
GST_RATES = (0, 3, 5, 12, 18, 28)


def _rupees(paise):
    if paise is None:
        return None
    return float(Decimal(paise) / 100)


def export_workbook(products):
    """Every product, one row each, as an .xlsx file's bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(HEADERS)
    for p in products:
        ws.append([
            p.sku, p.name, p.brand or "", p.subtitle or "",
            p.category.name if p.category_id else "",
            p.subcategory.name if p.subcategory_id else "",
            _rupees(p.price_paise), _rupees(p.mrp_paise),
            p.stock, p.weight_grams, p.tax_rate_bps / 100,
            "yes" if p.featured else "no", "yes" if p.active else "no",
            p.image_url or "", "\n".join(p.gallery or []), p.description or "",
            "\n".join(f"{i.get('q', '')} | {i.get('a', '')}" for i in (p.faq or []) if isinstance(i, dict)),
            p.slug or "", p.seo_title or "", p.seo_description or "", str(p.id),
        ])
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="FDE7D3")
    for cell in ws[1]:
        cell.font = bold
        cell.fill = fill
    ws.freeze_panes = "B2"
    widths = {"sku": 14, "name": 34, "description": 60, "gallery": 50, "faq": 60, "cover_url": 40, "id": 38}
    for i, header in enumerate(HEADERS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(header, 16)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    guide = wb.create_sheet("How to fill")
    guide.append(["Column", "What goes in it"])
    for header, help_text in COLUMNS:
        guide.append([header, help_text])
    guide.append([])
    guide.append(["Rules", "The SKU is the key. A row is checked in full before anything is saved; one bad row and nothing is saved, with every problem listed."])
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 110
    for cell in guide[1]:
        cell.font = bold

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _yes(value, default):
    text = _text(value).lower()
    if not text:
        return default, None
    if text in ("yes", "y", "true", "1"):
        return True, None
    if text in ("no", "n", "false", "0"):
        return False, None
    return default, f"“{text}” is not yes or no"


def _money(value, label, required):
    text = _text(value).replace(",", "").replace("₹", "")
    if not text:
        return None, (f"{label} is required" if required else None)
    try:
        amount = Decimal(text)
    except InvalidOperation:
        return None, f"{label} “{text}” is not a number"
    if amount < 0:
        return None, f"{label} cannot be negative"
    return int((amount * 100).to_integral_value()), None


def _whole(value, label, required=True):
    text = _text(value)
    if not text:
        return None, (f"{label} is required" if required else None)
    try:
        number = Decimal(text)
    except InvalidOperation:
        return None, f"{label} “{text}” is not a number"
    if number != number.to_integral_value() or number < 0:
        return None, f"{label} must be a whole number, 0 or more"
    return int(number), None


def read_workbook(data):
    """Parse and check every row. Returns (plans, errors): `plans` is one
    dict per row ready to write, `errors` is [(row number, message)]. Writes
    nothing."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — any unreadable file is one answer
        return [], [(0, "That is not an .xlsx file. Download the sheet, edit it, and upload it back.")]
    ws = wb["Products"] if "Products" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], [(0, "The sheet is empty.")]
    header = [_text(h).lower() for h in rows[0]]
    missing = [h for h in ("sku", "name", "category", "price_rupees", "stock", "weight_grams") if h not in header]
    if missing:
        return [], [(1, f"Missing columns: {', '.join(missing)}. Start from the downloaded sheet.")]
    col = {h: i for i, h in enumerate(header) if h}

    categories = {c.name.lower(): c for c in ShopCategory.objects.all()}
    subcategories = {(s.category_id, s.name.lower()): s for s in ShopSubcategory.objects.all()}
    existing = {p.sku: p for p in Product.objects.all()}
    by_id = {str(p.id): p for p in existing.values()}
    slugs_taken = {p.slug: p.sku for p in existing.values() if p.slug}

    plans, errors, seen_skus, seen_slugs = [], [], {}, {}
    for n, raw in enumerate(rows[1:], start=2):
        get = lambda h: raw[col[h]] if h in col and col[h] < len(raw) else None  # noqa: E731
        if all(_text(v) == "" for v in raw):
            continue
        problems = []
        sku = _text(get("sku")).upper()
        if not sku:
            problems.append("SKU is required — it is how a row finds its product")
        elif sku in seen_skus:
            problems.append(f"SKU {sku} is also on row {seen_skus[sku]}; one product, one row")
        else:
            seen_skus[sku] = n
        product = existing.get(sku)
        row_id = _text(get("id"))
        if row_id and row_id in by_id and by_id[row_id].sku != sku:
            problems.append(
                f"this row's id belongs to {by_id[row_id].sku}; to change a SKU, change it on that product in the console"
            )

        name = _text(get("name"))
        if not name:
            problems.append("name is required")
        category = categories.get(_text(get("category")).lower())
        if category is None:
            problems.append(
                f"category “{_text(get('category'))}” does not exist — one of: {', '.join(c.name for c in categories.values())}"
            )
        subcategory = None
        sub_name = _text(get("subcategory"))
        if sub_name and category is not None:
            subcategory = subcategories.get((category.id, sub_name.lower()))
            if subcategory is None:
                problems.append(f"subcategory “{sub_name}” is not under {category.name}")

        price, e1 = _money(get("price_rupees"), "price", True)
        mrp, e2 = _money(get("mrp_rupees"), "MRP", False)
        stock, e3 = _whole(get("stock"), "stock")
        weight, e4 = _whole(get("weight_grams"), "weight")
        problems += [e for e in (e1, e2, e3, e4) if e]
        if price is not None and mrp is not None and mrp <= price:
            problems.append("MRP must be above the price, or blank")
        if weight is not None and weight <= 0:
            problems.append("weight must be above 0 grams")

        gst_text = _text(get("gst_percent")) or "0"
        try:
            gst = Decimal(gst_text.replace("%", ""))
        except InvalidOperation:
            gst = None
        if gst is None or gst not in GST_RATES:
            problems.append(f"GST “{gst_text}” is not one of {', '.join(map(str, GST_RATES))}")
        featured, e5 = _yes(get("featured"), product.featured if product else False)
        active, e6 = _yes(get("active"), product.active if product else True)
        problems += [e for e in (e5, e6) if e]

        urls = []
        for label, value in (("cover_url", get("cover_url")), ("gallery", get("gallery"))):
            for line in _text(value).splitlines():
                line = line.strip()
                if line and not line.startswith(("https://", "http://")):
                    problems.append(f"{label}: “{line[:50]}” is not a web address")
                elif line:
                    urls.append((label, line))
        cover = next((u for label, u in urls if label == "cover_url"), "")
        gallery = [u for label, u in urls if label == "gallery"]

        faq = []
        for line in _text(get("faq")).splitlines():
            if not line.strip():
                continue
            if "|" not in line:
                problems.append(f"FAQ line “{line[:50]}” needs a | between question and answer")
                continue
            q, a = (part.strip() for part in line.split("|", 1))
            if not q or not a:
                problems.append(f"FAQ line “{line[:50]}” is missing its question or its answer")
            else:
                faq.append({"q": q, "a": a})

        typed_slug = slugify(_text(get("slug")))[:110]
        slug = typed_slug or (product.slug if product else None)
        if slug:
            owner = slugs_taken.get(slug)
            if owner and owner != sku:
                problems.append(f"slug “{slug}” is already {owner}'s page")
            elif slug in seen_slugs and seen_slugs[slug] != n:
                problems.append(f"slug “{slug}” is also on row {seen_slugs[slug]}")
            seen_slugs[slug] = n

        if problems:
            errors += [(n, f"{sku or 'no SKU'}: {p}") for p in problems]
            continue
        plans.append({
            "row": n, "sku": sku, "product": product,
            "fields": {
                "name": name, "brand": _text(get("brand")) or None,
                "subtitle": _text(get("subtitle")) or None,
                "category": category, "subcategory": subcategory,
                "price_paise": price, "mrp_paise": mrp, "stock": stock,
                "weight_grams": weight, "tax_rate_bps": int(gst * 100),
                "featured": featured, "active": active,
                "image_url": cover or None, "gallery": gallery,
                "description": _text(get("description")) or None, "faq": faq,
                "slug": slug, "seo_title": _text(get("seo_title")) or None,
                "seo_description": _text(get("seo_description")) or None,
            },
        })
    return plans, errors


def apply(plans):
    """Write checked rows in one transaction. Returns (added, updated,
    changes): `changes` lists, per product, the fields whose value moved —
    what the audit log records."""
    added, updated, changes = 0, 0, []
    with transaction.atomic():
        for plan in plans:
            product = plan["product"]
            if product is None:
                product = Product(id=uuid.uuid4(), sku=plan["sku"], created_at=timezone.now())
                for field, value in plan["fields"].items():
                    setattr(product, field, value)
                product.save()
                added += 1
                changes.append({"sku": plan["sku"], "added": True})
                continue
            moved = {}
            for field, value in plan["fields"].items():
                before = getattr(product, field)
                if before != value:
                    moved[field] = [_plain(before), _plain(value)]
                    setattr(product, field, value)
            if moved:
                product.save()
                updated += 1
                changes.append({"sku": plan["sku"], "changed": moved})
    return added, updated, changes


def _plain(value):
    """A value as the audit log can store it."""
    if hasattr(value, "name") and hasattr(value, "pk"):
        return value.name
    return value
